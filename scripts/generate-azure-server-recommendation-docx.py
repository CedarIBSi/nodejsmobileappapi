from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "IBS_Mobile_API_Azure_Server_Recommendation.docx"


def bullets(document: Document, values: list[str]) -> None:
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def numbered(document: Document, values: list[str]) -> None:
    for value in values:
        document.add_paragraph(value, style="List Number")


def table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    item = document.add_table(rows=1, cols=len(headers))
    item.style = "Table Grid"
    for index, header in enumerate(headers):
        item.rows[0].cells[index].text = header
    for row in rows:
        cells = item.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10.5)

title = doc.add_heading("IBS Intelligence Mobile API", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Azure Server Recommendation and IT Implementation Brief")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
date = doc.add_paragraph("Prepared: 4 September 2026")
date.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("1. Decision summary", level=1)
doc.add_paragraph(
    "Purchase one Azure Standard_D8as_v6 virtual machine with Windows Server 2025 Datacenter: Azure Edition, "
    "8 vCPUs, and 32 GB RAM. This server will host IIS, the Node.js API, PostgreSQL, "
    "and protected journal/white-paper files. The public WordPress website remains on Kinsta."
)
doc.add_paragraph(
    "If Standard_D8as_v6 is unavailable in the selected Azure region, use "
    "Standard_D8as_v5 with the same 8-vCPU/32-GB profile."
)

doc.add_heading("2. Business context", level=1)
bullets(doc, [
    "IBS Intelligence has an established website audience built over multiple years.",
    "The website analytics screenshot showed approximately 769,449 views and 603,109 active users for the displayed 2026 period.",
    "The website stays on Kinsta; the Azure VM is dedicated to the mobile backend, its database, and protected documents.",
    "The mobile launch may create short traffic spikes when promoted to the existing audience.",
    "The target is approximately 100,000 registered mobile users, not 100,000 simultaneous users.",
])

doc.add_heading("3. Required Azure VM", level=1)
table(doc, ["Setting", "Required selection"], [
    ["VM family/size", "Standard_D8as_v6"],
    ["Fallback size", "Standard_D8as_v5 if v6 is unavailable"],
    ["vCPU", "8"],
    ["Memory", "32 GB"],
    ["Operating system", "Windows Server 2025 Datacenter: Azure Edition (Desktop Experience), or Windows Server 2022 if required by IT policy"],
    ["Architecture", "x64"],
    ["Region", "Central India or South India, based on user proximity and service availability"],
    ["Availability", "Deploy in an Availability Zone"],
    ["Public IP", "Standard, static IPv4"],
    ["Security type", "Trusted Launch"],
    ["Secure Boot", "Enabled"],
    ["Virtual TPM", "Enabled"],
])

doc.add_heading("4. Required managed disks", level=1)
table(doc, ["Disk", "Capacity", "Type", "Purpose"], [
    ["OS disk", "128 GB", "Premium SSD", "Windows Server, application code, Node.js, IIS, packages, and system logs"],
    ["Database disk", "256 GB", "Premium SSD v2", "PostgreSQL data; begin near 5,000 IOPS and 200–250 MB/s"],
    ["Document disk", "512 GB", "Premium SSD", "Protected journal and white-paper PDFs"],
])
doc.add_paragraph(
    "PostgreSQL and PDFs must not use Azure temporary/local storage. Separate managed disks "
    "prevent PDF downloads and operating-system activity from competing directly with database I/O."
)

doc.add_heading("5. Windows storage layout", level=1)
doc.add_paragraph(
    "C:\\IBS\\API\n"
    "D:\\PostgreSQL\\Data\n"
    "E:\\IBSI-PDFs\\Journals\n"
    "E:\\IBSI-PDFs\\Whitepapers\n"
    "C:\\inetpub\\logs\n"
    "C:\\IBS\\Logs"
)
doc.add_paragraph("Application environment settings:")
doc.add_paragraph(
    "JOURNAL_STORAGE_DIR=E:\\IBSI-PDFs\\Journals\n"
    "WHITEPAPER_STORAGE_DIR=E:\\IBSI-PDFs\\Whitepapers"
)

doc.add_heading("6. Network and firewall requirements", level=1)
table(doc, ["Port", "Source", "Rule"], [
    ["443/TCP", "Internet", "Allow HTTPS"],
    ["80/TCP", "Internet", "Allow only HTTP-to-HTTPS redirect and certificate validation"],
    ["3389/TCP", "Office static IP, VPN, or Azure Bastion only", "Allow restricted RDP administration"],
    ["3000/TCP", "Internet", "Deny; Node must listen on localhost"],
    ["5432/TCP", "Internet", "Deny; PostgreSQL must listen on localhost"],
    ["All other inbound traffic", "Internet", "Deny"],
])
bullets(doc, [
    "Node.js should listen on 127.0.0.1:3000.",
    "PostgreSQL should listen on 127.0.0.1:5432.",
    "IIS should be the only public application entry point on ports 80 and 443.",
    "Use a Cloudflare-proxied DNS record/WAF in front of Azure where possible.",
    "Do not expose RDP to the general Internet; use Azure Bastion or a restricted VPN/source IP with MFA-enabled administrator access.",
])

doc.add_heading("7. Software to install", level=1)
bullets(doc, [
    "Current Windows Server security and cumulative updates.",
    "Current Node.js LTS release supported by the application.",
    "IIS with URL Rewrite and Application Request Routing (ARR) for HTTPS termination and reverse proxying.",
    "PostgreSQL 18 current minor release, or another approved supported PostgreSQL release.",
    "Use the Node.js PostgreSQL driver's built-in connection pools initially; add an approved external pooler later if measurements justify it.",
    "Git or a controlled deployment package mechanism.",
    "Azure Monitor Agent and required backup agent/extensions.",
])

doc.add_heading("8. Application process design", level=1)
bullets(doc, [
    "Run three or four Node.js worker processes behind IIS ARR.",
    "Install each worker as a Windows service using WinSW or an IT-approved service wrapper, with automatic startup, restart-on-failure, and graceful shutdown.",
    "Do not run production with an interactive npm start terminal.",
    "Set NODE_ENV=production and LOG_LEVEL=info.",
    "Begin with a PostgreSQL pool of 5–10 connections per Node worker.",
    "Keep the total database pool size controlled across all workers and reassess external connection pooling before increasing worker counts significantly.",
])

doc.add_heading("9. IIS reverse-proxy requirements", level=1)
bullets(doc, [
    "Install IIS, URL Rewrite, and Application Request Routing; proxy only to Node.js workers listening on localhost.",
    "Enable supported TLS protocols, HTTP/2, upstream keep-alive, and dynamic JSON compression.",
    "Configure request, connect, send, and upstream response timeouts.",
    "Apply rate limiting at Cloudflare/Azure WAF as well as inside the API.",
    "Do not log signed journal or white-paper viewing tokens.",
    "Configure IIS access/error log rotation and retention.",
    "Automate TLS certificate issuance and renewal, with expiry alerts.",
])

doc.add_heading("10. PostgreSQL starting configuration", level=1)
doc.add_paragraph("Initial values for a combined 32-GB server, subject to measurement and IT review:")
doc.add_paragraph(
    "shared_buffers = 8GB\n"
    "effective_cache_size = 20GB\n"
    "work_mem = 8MB\n"
    "maintenance_work_mem = 1GB\n"
    "max_connections = 100\n"
    "random_page_cost = 1.1\n"
    "effective_io_concurrency = 200\n"
    "log_min_duration_statement = 500ms"
)
bullets(doc, [
    "Enable pg_stat_statements.",
    "Monitor autovacuum, table/index growth, slow queries, locks, and connections.",
    "Keep PostgreSQL inaccessible from the public Internet.",
    "Tune final values from observed production memory and query behavior rather than applying them blindly.",
])

doc.add_heading("11. Secrets and security", level=1)
bullets(doc, [
    "Store Firebase, Google Play, database, and signing credentials in Azure Key Vault where practical.",
    "If an environment file is temporarily used, restrict its NTFS permissions to Administrators and the dedicated API service identity.",
    "Do not commit credentials, private keys, service-account JSON, TLS keys, or production .env files to Git.",
    "Use managed identities wherever supported.",
    "Enable Microsoft Defender for Cloud and automatic OS security updates.",
    "Rotate credentials during migration if the existing secrets may have been readable by other local users.",
])

doc.add_heading("12. Backup and recovery", level=1)
bullets(doc, [
    "Enable Azure Backup for the complete VM using a Recovery Services vault.",
    "Use geo-redundant backup storage where budget and policy permit.",
    "Create daily PostgreSQL logical backups in addition to VM backup.",
    "Configure PostgreSQL WAL archiving/point-in-time recovery if the required recovery point is shorter than one day.",
    "Back up the PDF managed disk separately.",
    "Keep at least 30 daily recovery points and an agreed monthly long-term recovery point.",
    "Perform and document a full restoration test at least quarterly.",
])

doc.add_heading("13. Monitoring and alerts", level=1)
table(doc, ["Metric", "Initial alert threshold"], [
    ["CPU", "Above 75% for 10 minutes"],
    ["Memory", "Above 80%"],
    ["Disk usage", "Above 75%"],
    ["Disk latency", "Sustained above 10–20 ms"],
    ["PostgreSQL connections", "Above 70"],
    ["API p95 response time", "Above 500 ms"],
    ["HTTP 5xx rate", "Above 1%"],
    ["TLS certificate", "Alert well before expiry"],
    ["Backups", "Alert on any failed or missed backup"],
    ["Node/PostgreSQL/IIS", "Alert when any service stops or repeatedly restarts"],
])

doc.add_heading("14. Migration checklist", level=1)
numbered(doc, [
    "Create the Azure VM, managed disks, virtual network, Network Security Group, static IP, and backup vault.",
    "Harden Windows Server, apply updates, and create a dedicated least-privilege API service account.",
    "Attach and format the managed disks as NTFS volumes, assign fixed D: and E: drive letters, and restrict their ACLs.",
    "Install and configure PostgreSQL, Node.js, IIS, URL Rewrite, and ARR.",
    "Copy the API into C:\\IBS\\API and install production dependencies.",
    "Migrate PostgreSQL using a verified backup/restore procedure.",
    "Copy journals and white papers while preserving filenames and permissions.",
    "Install production secrets through Key Vault or a restricted environment file.",
    "Run all migrations in order and verify the migration ledger.",
    "Build the API and install its Windows service workers using WinSW or the IT-approved service manager.",
    "Configure IIS ARR, TLS, log redaction, Cloudflare/Azure WAF rate limits, and upstream health checks.",
    "Test Firebase sign-in/profile synchronization, subscriptions, entitlements, PDFs, range requests, and public endpoints.",
    "Test backups and perform at least one restoration before DNS cutover.",
    "Reduce DNS TTL, perform cutover, monitor errors/latency, and retain a documented rollback path.",
])

doc.add_heading("15. Acceptance criteria", level=1)
bullets(doc, [
    "Only ports 80 and 443 are publicly reachable; administration is restricted.",
    "The API and PostgreSQL automatically start after a VM reboot.",
    "The health endpoint reports HTTP 200 and database OK.",
    "Firebase profile synchronization succeeds with the production mobile application.",
    "Unauthorized users cannot access journal or white-paper links.",
    "Signed PDF links expire and their bearer tokens are absent from IIS/application logs.",
    "Google Play and Apple purchase verification behave correctly for configured platforms.",
    "Automated backup completes and a restoration test succeeds.",
    "Monitoring and alert notifications reach the responsible IT team.",
    "A staged load test is completed before a large marketing campaign.",
])

doc.add_heading("16. Scaling guidance", level=1)
doc.add_paragraph(
    "Start with Standard_D8as_v6 and monitor the first three months. Resize only when CPU "
    "regularly exceeds 70%, memory exceeds 80%, disk latency is sustained, or tested traffic "
    "cannot meet the response-time objective. A larger single VM improves capacity but does not "
    "provide high availability. The preferred later upgrade is to separate PostgreSQL into Azure "
    "Database for PostgreSQL, run two or more API instances, and move PDFs to private object storage."
)

doc.add_heading("17. IT review points", level=1)
bullets(doc, [
    "Confirm Central India versus South India based on users, company policy, availability zones, and VM/disk availability.",
    "Confirm expected PDF collection size and annual growth before finalizing the 512-GB document disk.",
    "Confirm recovery-point and recovery-time objectives and resulting backup retention.",
    "Confirm whether Cloudflare, Azure Front Door, or Application Gateway will terminate/protect Internet traffic.",
    "Confirm whether Azure Key Vault and Azure Bastion are approved and included in the budget.",
    "Obtain an Azure cost estimate for compute, three disks, public IP, outbound data, backups, monitoring, and security services.",
])

doc.add_heading("18. Final recommendation", level=1)
final = doc.add_paragraph()
final.add_run(
    "Approve one Standard_D8as_v6 Windows Server 2025 Datacenter: Azure Edition VM with 8 vCPUs, 32 GB RAM, a 128-GB Premium SSD "
    "OS disk, a 256-GB Premium SSD v2 database disk, and a 512-GB Premium SSD document disk."
).bold = True
final.add_run(
    " Deploy it in an Indian Azure region and availability zone, protect it with strict network "
    "rules and Cloudflare/WAF, enable Azure Backup and monitoring, and reassess architecture after "
    "real mobile traffic is measured."
)

doc.save(OUTPUT)
print(OUTPUT)
