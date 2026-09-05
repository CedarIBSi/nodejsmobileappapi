from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "IBS_API_Security_and_Scalability_Audit.docx"


def bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)

title = doc.add_heading("IBS Intelligence API", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Security, Scalability, and 100K-User Readiness Audit")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
date = doc.add_paragraph("Assessment date: 2 September 2026")
date.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Executive verdict", level=1)
doc.add_paragraph(
    "The API has a solid security foundation, and PostgreSQL can comfortably store "
    "100,000 users. The current single Windows server is suitable for development, "
    "testing, and a limited launch, but it is not yet ready to promise heavy 100K-user "
    "traffic, high availability, or super-speed performance. The primary limitations "
    "are architecture and operations rather than the database schema."
)

doc.add_heading("What is already good", level=1)
bullets(doc, [
    "Firebase authentication with revoked-token checking.",
    "Parameterized SQL and no obvious SQL-injection path found.",
    "Zod validation on request bodies, query parameters, and path parameters.",
    "Helmet security headers, rate limiting, and one-megabyte request limits.",
    "Google and Apple purchase states are verified server-to-server.",
    "Cryptographically verified webhooks and timing-safe signed PDF tokens.",
    "PDF path-traversal protections and entitlement checks for private content.",
    "Appropriate database indexes for the current query patterns.",
    "Strict TypeScript checking passes; Firebase, PostgreSQL, and HTTPS are operational.",
])

doc.add_heading("Priority security findings", level=1)
security_rows = [
    ["High", "Secret-file permissions", ".env and Firebase JSON are readable by BUILTIN\\Users.", "Restrict ACLs to the API service identity, SYSTEM, and Administrators; rotate credentials if exposure is possible."],
    ["High", "Signed PDF tokens in Nginx logs", "Nginx logs the full journal/white-paper bearer URL.", "Disable or redact access logging for signed viewing paths; prefer object-storage signed URLs."],
    ["High", "Account deletion and billing", "Deletion checks omit current store states such as trialing and grace-period states.", "Base deletion eligibility on active entitlements and authoritative store state."],
    ["Moderate", "Vulnerable qs dependency", "qs 6.15.3 has denial-of-service advisories; npm reports a fix.", "Apply the audited dependency update and regression-test the API."],
    ["Moderate", "Open CORS", "Blank CORS_ORIGINS allows arbitrary origins with credentials enabled.", "Configure an explicit production allowlist."],
    ["Moderate", "Broad network bindings", "Node 3000 and PostgreSQL 5432 bind to all interfaces.", "Bind both to localhost and expose only Nginx 80/443."],
    ["Moderate", "Notification authorization", "All staff roles, including employee, can notify any user.", "Restrict notification sending to admin or super_admin."],
    ["Moderate", "Unbounded memory cache", "Search/page/article cache entries expire logically but are not proactively evicted.", "Use Redis TTLs or a bounded LRU cache."],
]
add_table(doc, ["Severity", "Finding", "Risk", "Required action"], security_rows)

doc.add_heading("Current server inventory", level=1)
add_table(doc, ["Component", "Observed configuration"], [
    ["Operating system", "Windows Server 2019 Standard, build 17763.8755"],
    ["CPU", "6 physical cores / 12 logical processors, Intel Xeon Silver 4310"],
    ["Memory", "16 GB RAM; approximately 10.6 GB free during inspection"],
    ["Disk", "Approximately 199 GB total / 145 GB free"],
    ["Node.js", "24.15.0, one API process"],
    ["Nginx", "1.31.3 for Windows, one effective worker, 1,024 worker connections"],
    ["PostgreSQL", "18.4, local database approximately 9.6 MB"],
    ["Node database pool", "Maximum 10 PostgreSQL connections"],
    ["PostgreSQL", "100 max connections, 128 MB shared_buffers, 4 GB effective_cache_size"],
])

doc.add_heading("Can it handle 100K users?", level=1)
doc.add_paragraph(
    "100,000 registered users are not a problem for PostgreSQL. One hundred thousand "
    "simultaneous users are far beyond the current deployment. A practical planning "
    "model of 100,000 daily users making ten API calls each produces approximately one "
    "million calls per day: about 12 requests per second on average and roughly 120–240 "
    "requests per second during a 10–20x peak. The hardware could support that class of "
    "workload after architectural tuning, but the current deployment cannot guarantee it."
)

doc.add_heading("Primary scale bottlenecks", level=1)
add_table(doc, ["Area", "Current limitation", "Recommended direction"], [
    ["Nginx on Windows", "Only one effective worker; official documentation says high scalability should not be expected.", "Move Nginx and Node to Linux."],
    ["Node runtime", "One process and one event loop; no confirmed supervisor/autorestart.", "Run 2–4 supervised API instances behind a load balancer."],
    ["Firebase", "Revocation checking can add a Firebase network call to every private request.", "Measure latency; reserve uncached revocation checks for sensitive actions or use a carefully designed short cache/session strategy."],
    ["PDF delivery", "Every PDF byte and range request passes through Node and local disk.", "Move PDFs to object storage and a CDN using short-lived signed URLs."],
    ["WordPress", "Cloudflare currently returns 403, causing API 502 responses.", "Allowlist the backend or configure a secure WAF bypass, then cache at the CDN/Redis layer."],
    ["Cache/rate limit", "Both are process-local and unsuitable for multiple API instances.", "Use Redis for shared caching and distributed rate limiting."],
    ["Database", "Database shares the application host and uses conservative defaults.", "Use PgBouncer and a dedicated/managed PostgreSQL deployment with backups and monitoring."],
])

doc.add_heading("Database growth", level=1)
doc.add_paragraph(
    "At five metered articles per month, 100,000 users can create 500,000 access rows per "
    "month or approximately six million rows per year. PostgreSQL can handle this with the "
    "existing indexing approach, but production needs retention policies, partitioning for "
    "high-growth tables, autovacuum monitoring, and slow-query monitoring."
)
bullets(doc, [
    "Consider monthly partitioning for news_article_access and store_events.",
    "Enable pg_stat_statements and log queries slower than approximately 250–500 ms.",
    "Tune shared_buffers and effective_cache_size after separating application and database workloads.",
    "Use PgBouncer before adding many Node instances so connection counts remain bounded.",
    "Create encrypted daily backups, point-in-time recovery, and routine restoration tests.",
])

doc.add_heading("Operational findings", level=1)
bullets(doc, [
    "npm run build encounters Windows EPERM errors on some existing dist files, although no-emit TypeScript checking passes.",
    "Migration 012_drop_razorpay.sql is not recorded as applied while migration 013 is applied; repair ordering after a verified backup.",
    "No automated test suite is present.",
    "No confirmed process supervisor/autorestart configuration was found for the API.",
    "No confirmed automated PostgreSQL backup task was found; verify whether backups are managed externally.",
    "No confirmed Nginx log-rotation automation was found.",
    "Google Pub/Sub webhook configuration and Apple purchase verification are incomplete.",
    "Google Play plans still use test product identifiers.",
])

doc.add_heading("Software maintenance", level=1)
add_table(doc, ["Software", "Installed", "Observed current release", "Action"], [
    ["Node.js", "24.15.0", "24.20.0 LTS", "Upgrade through a tested maintenance window."],
    ["Nginx", "1.31.3", "1.31.4 mainline", "Upgrade; longer-term move from Windows to Linux."],
    ["PostgreSQL", "18.4", "18.6", "Apply the current supported minor release after backup/testing."],
    ["Windows Server", "2019 build 17763.8755", "August 2026 build 17763.9121", "Apply current security updates; extended support ends in January 2029."],
])

doc.add_heading("Recommended target architecture", level=1)
architecture = doc.add_paragraph()
architecture.alignment = WD_ALIGN_PARAGRAPH.CENTER
architecture.add_run(
    "Cloudflare / WAF / CDN\n"
    "↓\n"
    "Linux load balancer\n"
    "↓\n"
    "2–4 Node API instances\n"
    "↓\n"
    "Redis cache and distributed rate limiter\n"
    "↓\n"
    "PgBouncer\n"
    "↓\n"
    "Managed or dedicated PostgreSQL\n\n"
    "PDFs → object storage + CDN signed URLs\n"
    "Logs and metrics → centralized monitoring and alerts"
).bold = True

doc.add_heading("Implementation order", level=1)
doc.add_heading("Before public launch", level=2)
numbered(doc, [
    "Fix secret-file ACLs and rotate credentials where necessary.",
    "Stop Nginx from logging signed PDF tokens.",
    "Fix account-deletion subscription checks.",
    "Apply dependency security updates and add regression tests.",
    "Fix WordPress/Cloudflare connectivity.",
    "Repair migration 012/013 consistency after a database backup.",
    "Finish Apple and Google Pub/Sub configuration.",
    "Replace test store product IDs.",
    "Establish encrypted backups and prove restoration works.",
    "Fix repeatable production builds and configure automatic process restart.",
])

doc.add_heading("Before a large campaign", level=2)
numbered(doc, [
    "Move Nginx and Node to Linux.",
    "Run multiple supervised Node API instances.",
    "Add Redis caching and distributed rate limiting.",
    "Move PDFs to object storage/CDN.",
    "Tune PostgreSQL and add PgBouncer.",
    "Add uptime, error, CPU, memory, disk, queue, and slow-query monitoring.",
    "Add unit, integration, authorization, webhook, and security tests.",
    "Run staged load tests at 50, 100, 250, 500, and 1,000 requests per second.",
])

doc.add_heading("Official references", level=1)
bullets(doc, [
    "Nginx for Windows limitations: https://nginx.org/en/docs/windows.html",
    "Nginx releases and security news: https://nginx.org/2026.html",
    "Node.js release schedule: https://nodejs.org/en/about/previous-releases",
    "PostgreSQL versioning policy: https://www.postgresql.org/support/versioning/",
    "Windows Server release information: https://learn.microsoft.com/en-us/windows/release-health/windows-server-release-info",
])

doc.add_heading("Final conclusion", level=1)
doc.add_paragraph(
    "The API design is promising and its database model can support 100,000 users. The "
    "current Windows single-server deployment needs targeted security fixes, reliable "
    "operations, shared infrastructure, and horizontal scaling before it can safely promise "
    "high concurrency, high availability, or super-speed performance."
)

doc.save(OUTPUT)
print(OUTPUT)
